"""Five-year financial model of proposals 1-3 (Financial impact slide, its appendix table, the Excel model).

    python storyboard/fin_model.py        # writes storyboard/chagee_financial_model.xlsx and prints the yearly rows

compute() and the workbook use the same inputs from assumptions.json (fin_model plus the deck's base volumes, price,
margin, wage and machine cost) and the same arithmetic: the workbook holds formulas, compute() mirrors them for the
slide. The Model sheet is incremental to the business without the proposals; the SG forecast sheet is the whole of Chagee
Singapore, with and without them. Every figure is illustrative. Nothing is typed here.
"""
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
RAW = json.loads((ROOT / "assumptions.json").read_text(encoding="utf-8"))
OUT = HERE / "chagee_financial_model.xlsx"


def inputs(raw: dict = RAW) -> dict:
    F = {k: o["value"] for k, o in raw["fin_model"].items() if not k.startswith("_")}
    F.update(outlets=raw["facts"]["sg_outlets"]["value"], cups=F["cups_per_outlet"],
             price=raw["storyboard"]["price_per_cup"]["value"], margin=F["gross_margin"],
             wage=raw["costs"]["labour_cost_per_hour"]["value"], machine=raw["costs"]["machine_cost_per_unit_per_month"]["value"],
             open_hours=raw["store"]["open_hours_per_day"]["value"], open_days=raw["store"]["open_days_per_year"]["value"],
             rent=raw["site_ext"]["rent_per_month"]["value"], utilities=raw["site_ext"]["utilities_per_month"]["value"],
             delivery_share=raw["site_ext"]["delivery_share_of_cups"]["value"], delivery_commission=raw["site_ext"]["delivery_commission_share"]["value"])
    return F


def compute(raw: dict = RAW) -> dict:
    """Yearly rows, same names and arithmetic as the Model sheet."""
    F = inputs(raw)
    n = len(F["years"]); R = {"years": F["years"]}
    R["growth"] = [(1 + F["organic_growth"]) ** i for i in range(n)]                          # same-store growth index
    R["wage_idx"] = [(1 + F["wage_inflation"]) ** i for i in range(n)]
    opened = [sum(F["new_stores"][:i + 1]) - F["new_stores"][i] / 2 for i in range(n)]      # store-years: half a year when opened
    R["store_years"] = opened
    R["outlets_avg"] = [F["outlets"] + opened[i] for i in range(n)]
    R["outlets_end"] = [F["outlets"] + sum(F["new_stores"][:i + 1]) for i in range(n)]
    R["base_cups"] = [R["outlets_avg"][i] * F["cups"] * R["growth"][i] for i in range(n)]
    R["base_rev"] = [c * F["price"] for c in R["base_cups"]]
    base_cups, base_rev = R["base_cups"], R["base_rev"]
    p1 = [F["p1_coverage"][i] * F["effectiveness"][i] for i in range(n)]
    p2 = [F["p2_coverage"][i] * F["effectiveness"][i] for i in range(n)]
    R["p1_cups"] = [base_cups[i] * F["p1_lost_share"] * F["p1_recovery_share"] * p1[i] for i in range(n)]
    R["p2_cups"] = [base_cups[i] * F["p2_extra_cup_share"] * p2[i] for i in range(n)]
    R["p3_cups"] = [F["cups"] * F["p3_uplift"] * opened[i] * R["growth"][i] for i in range(n)]
    for k in ("p1", "p2", "p3"):
        R[f"{k}_rev"] = [c * F["price"] for c in R[f"{k}_cups"]]
    R["cups"] = [R["p1_cups"][i] + R["p2_cups"][i] + R["p3_cups"][i] for i in range(n)]
    R["rev"] = [R["p1_rev"][i] + R["p2_rev"][i] + R["p3_rev"][i] for i in range(n)]
    R["cogs"] = [-r * (1 - F["margin"]) for r in R["rev"]]
    R["gp"] = [R["rev"][i] + R["cogs"][i] for i in range(n)]
    R["hours"] = [F["p2_idle_hours_week"] * 52 * R["outlets_avg"][i] * p2[i] for i in range(n)]
    R["labour"] = [R["hours"][i] * F["wage"] * R["wage_idx"][i] for i in range(n)]
    R["waste"] = [base_rev[i] * F["p2_waste_share"] * F["p2_waste_cut"] * p2[i] for i in range(n)]
    R["machines"] = [R["outlets_avg"][i] * F["p2_machines_saved"] * F["machine"] * 12 * p2[i] for i in range(n)]
    R["savings"] = [R["labour"][i] + R["waste"][i] + R["machines"][i] for i in range(n)]
    R["benefit"] = [R["gp"][i] + R["savings"][i] for i in range(n)]
    R["team"] = [m * (F["pm_cost_month"] + F["engineer_cost_month"]) for m in F["team_months"]]
    R["platform"] = [s * F["platform_build"] for s in F["build_share"]]
    R["data"] = [s * F["data_setup"] for s in F["build_share"]]
    R["training"] = [o * F["training_per_outlet"] for o in F["training_outlets"]]
    R["security"] = [s * F["security_review"] for s in F["build_share"]]
    R["fixed"] = [R["team"][i] + R["platform"][i] + R["data"][i] + R["security"][i] + R["training"][i] for i in range(n)]
    R["run"] = [s * F["run_team_year"] for s in F["run_share"]]
    R["maint"] = [s * F["maintenance_share"] * F["platform_build"] for s in F["run_share"]]
    R["cloud"] = [s * F["cloud_year"] for s in F["cloud_share"]]
    R["chagee"] = list(F["chagee_time_year"])
    R["perks"] = [base_cups[i] * F["p1_redirect_share"] * p1[i] * F["p1_voucher_share"] * F["p1_voucher_cost"] for i in range(n)]
    R["variable"] = [R["run"][i] + R["maint"][i] + R["cloud"][i] + R["chagee"][i] + R["perks"][i] for i in range(n)]
    R["cost"] = [R["fixed"][i] + R["variable"][i] for i in range(n)]
    R["net"] = [R["benefit"][i] - R["cost"][i] for i in range(n)]
    R["cum"] = [sum(R["net"][:i + 1]) for i in range(n)]
    # the whole of Chagee Singapore, with the proposals (and the same business without them)
    oa = R["outlets_avg"]
    R["sg_rev"] = [base_rev[i] + R["rev"][i] for i in range(n)]
    R["sg_cups"] = [base_cups[i] + R["cups"][i] for i in range(n)]
    # an income statement: Proposal 2's savings sit in the lines they reduce
    R["sg_cogs"] = [-R["sg_rev"][i] * (1 - F["margin"]) + R["waste"][i] for i in range(n)]          # net of fresh-stock waste avoided
    R["sg_gp"] = [R["sg_rev"][i] + R["sg_cogs"][i] for i in range(n)]
    R["sg_gm"] = [R["sg_gp"][i] / R["sg_rev"][i] for i in range(n)]
    R["sg_staff_gross"] = [-oa[i] * F["crew_on_shift"] * F["open_hours"] * F["open_days"] * F["wage"] * R["wage_idx"][i] for i in range(n)]
    R["sg_labour"] = [R["sg_staff_gross"][i] + R["labour"][i] for i in range(n)]                      # net of idle hours saved
    R["sg_rent"] = [-oa[i] * (F["rent"] + F["utilities"]) * 12 for i in range(n)]
    R["sg_delivery"] = [-base_rev[i] * F["delivery_share"] * F["delivery_commission"] for i in range(n)]
    R["sg_equip"] = list(R["machines"])                                                              # tea machines avoided
    R["sg_ai"] = [-c for c in R["cost"]]
    R["sg_opex"] = [R["sg_labour"][i] + R["sg_rent"][i] + R["sg_delivery"][i] + R["sg_equip"][i] + R["sg_ai"][i] for i in range(n)]
    R["sg_op"] = [R["sg_gp"][i] + R["sg_opex"][i] for i in range(n)]
    R["sg_margin"] = [R["sg_op"][i] / R["sg_rev"][i] for i in range(n)]
    R["sg_op_base"] = [base_rev[i] * F["margin"] + R["sg_staff_gross"][i] + R["sg_rent"][i] + R["sg_delivery"][i] for i in range(n)]
    R["sg_margin_base"] = [R["sg_op_base"][i] / base_rev[i] for i in range(n)]
    R["sg_uplift"] = [R["sg_op"][i] - R["sg_op_base"][i] for i in range(n)]
    T = {k: sum(v) for k, v in R.items() if k not in ("years", "cum", "store_years", "growth", "wage_idx", "outlets_avg", "outlets_end", "sg_margin", "sg_margin_base", "sg_gm")}
    T["roi"] = T["net"] / T["cost"]
    T["payback"] = next((F["years"][i] for i in range(n) if R["cum"][i] > 0), None)
    T["cum"] = R["cum"][-1]
    return {"rows": R, "total": T, "inputs": F}


# --------------------------------------------------------------------------------------------- the Excel workbook
def export(path: Path = OUT) -> Path:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L

    F = inputs(); fm = RAW["fin_model"]; yrs = F["years"]; n = len(yrs)
    BLUE, BLACK, GREEN, RED = "0000FF", "000000", "008000", "B8322B"
    f = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)
    HEAD = PatternFill("solid", fgColor="24406B"); SEC = PatternFill("solid", fgColor="F3EFE7"); KEY = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="BFB6A6"); TOP = Border(top=thin); DBL = Border(top=thin, bottom=Side(style="double", color="BFB6A6"))
    SGD = '#,##0;[Red](#,##0);"-"'; NUM = '#,##0;(#,##0);"-"'; PCT = '0.0%;(0.0%);"-"'; DEC = '#,##0.00;(#,##0.00);"-"'
    wb = Workbook()

    # ---- Assumptions
    A = wb.active; A.title = "Assumptions"
    A["A1"] = "Chagee Singapore · AI proposals 1-3 · model assumptions"; A["A1"].font = f(bold=True, size=14)
    A["A2"] = ("All money in Singapore dollars (S$). Illustrative assumptions, not Chagee data. Blue cells are inputs: change them "
               "and the Model and Cost plan sheets update. Base figures reuse the inputs behind the deck's earlier slides.")
    A["A2"].font = f(italic=True, color="5C5C5C"); A.merge_cells("A2:H2"); A["A2"].alignment = Alignment(wrap_text=True); A.row_dimensions[2].height = 30
    ref = {}; r = 4

    def section(title, cols=("Input", "Value", "Unit", "Source / rationale")):
        nonlocal r
        A.cell(r, 1, title).font = f(bold=True, color="FFFFFF")
        for c in range(1, 9):
            A.cell(r, c).fill = HEAD
        for i, h in enumerate(cols[1:], 2):
            A.cell(r, i, h).font = f(bold=True, color="FFFFFF")
        r += 1

    def scalar(key, label, value, unit, note, fmt=NUM, key_cell=False):
        nonlocal r
        A.cell(r, 1, label).font = f()
        c = A.cell(r, 2, value); c.font = f(color=BLUE); c.number_format = fmt
        if key_cell:
            c.fill = KEY
        A.cell(r, 3, unit).font = f(color="5C5C5C")
        A.cell(r, 4, note).font = f(color="5C5C5C"); A.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
        ref[key] = f"Assumptions!$B${r}"; r += 1

    src = lambda sec, k: f"{RAW[sec][k]['source']}. {RAW[sec][k].get('rationale', '')}".strip()
    section("Base business (from the deck)")
    scalar("outlets", "Chagee outlets in Singapore", F["outlets"], "outlets", src("facts", "sg_outlets"))
    scalar("cups", "Cups per outlet per year (stabilised)", F["cups"], "cups", fm["cups_per_outlet"]["source"] + ".", NUM, True)
    scalar("price", "Average price per cup", F["price"], "S$ per cup", "Deck assumption (price per cup), illustrative.", DEC)
    scalar("margin", "Gross margin", F["margin"], "share of revenue", fm["gross_margin"]["source"] + ".", PCT, True)
    scalar("organic_growth", "Organic growth in cups per outlet", F["organic_growth"], "per year", fm["organic_growth"]["rationale"], PCT, True)
    scalar("wage", "Labour cost per hour", F["wage"], "S$ per hour", "Deck assumption (costs.labour_cost_per_hour), illustrative.", DEC)
    scalar("wage_inflation", "Wage inflation", F["wage_inflation"], "per year", fm["wage_inflation"]["rationale"], PCT, True)
    scalar("machine", "Tea machine cost per unit per month", F["machine"], "S$ per month", "Deck assumption (costs.machine_cost_per_unit_per_month), illustrative.", DEC)
    r += 1
    section("Singapore operations (SG forecast sheet only)")
    scalar("crew_on_shift", "Staff on shift, average", F["crew_on_shift"], "staff", fm["crew_on_shift"]["rationale"], DEC)
    scalar("open_hours", "Open hours per day", F["open_hours"], "hours", "Deck assumption (store.open_hours_per_day), illustrative.")
    scalar("open_days", "Open days per year", F["open_days"], "days", "Deck assumption (store.open_days_per_year), illustrative.")
    scalar("rent", "Rent per outlet", F["rent"], "S$ per month", "Deck assumption (site_ext.rent_per_month), illustrative.")
    scalar("utilities", "Utilities per outlet", F["utilities"], "S$ per month", "Deck assumption (site_ext.utilities_per_month), illustrative.")
    scalar("delivery_share", "Cups sold through delivery apps", F["delivery_share"], "share of baseline cups", "Deck assumption (site_ext.delivery_share_of_cups), illustrative.", PCT)
    scalar("delivery_commission", "Delivery-app commission", F["delivery_commission"], "share of delivery sales", "Deck assumption (site_ext.delivery_commission_share). Charged on baseline sales; the proposals' extra cups are mostly in-store and app pickups.", PCT)
    r += 1
    note = lambda k: fm[k]["rationale"]
    section("Proposal 1 · dynamic demand management")
    scalar("p1_lost_share", "Cups lost at peak today", F["p1_lost_share"], fm["p1_lost_share"]["unit"], note("p1_lost_share"), PCT, True)
    scalar("p1_recovery_share", "Share of lost cups recovered", F["p1_recovery_share"], fm["p1_recovery_share"]["unit"], note("p1_recovery_share"), PCT, True)
    scalar("p1_redirect_share", "Orders redirected to another outlet", F["p1_redirect_share"], fm["p1_redirect_share"]["unit"], note("p1_redirect_share"), PCT)
    scalar("p1_voucher_share", "Redirected orders given a switch perk", F["p1_voucher_share"], fm["p1_voucher_share"]["unit"], note("p1_voucher_share"), PCT)
    scalar("p1_voucher_cost", "Cost per switch perk", F["p1_voucher_cost"], "S$ per perk", note("p1_voucher_cost"), DEC)
    r += 1
    section("Proposal 2 · AI-enabled in-store planning")
    scalar("p2_extra_cup_share", "Extra cups served at the peak", F["p2_extra_cup_share"], fm["p2_extra_cup_share"]["unit"], note("p2_extra_cup_share"), PCT)
    scalar("p2_idle_hours_week", "Idle staff hours avoided", F["p2_idle_hours_week"], "hours per outlet per week", note("p2_idle_hours_week"), NUM, True)
    scalar("p2_waste_share", "Fresh stock wasted today", F["p2_waste_share"], fm["p2_waste_share"]["unit"], note("p2_waste_share"), PCT)
    scalar("p2_waste_cut", "Share of that waste avoided", F["p2_waste_cut"], fm["p2_waste_cut"]["unit"], note("p2_waste_cut"), PCT, True)
    scalar("p2_machines_saved", "Net machines avoided per outlet", F["p2_machines_saved"], "units per outlet", note("p2_machines_saved"), DEC)
    r += 1
    section("Proposal 3 · new-store site model")
    scalar("p3_uplift", "Extra cups at an AI-chosen site", F["p3_uplift"], fm["p3_uplift"]["unit"], note("p3_uplift"), PCT, True)
    r += 1
    section("Costs of the initiatives")
    scalar("pm_cost_month", "Project manager", F["pm_cost_month"], "S$ per month", note("pm_cost_month"))
    scalar("engineer_cost_month", "AI engineer", F["engineer_cost_month"], "S$ per month", note("engineer_cost_month"))
    scalar("platform_build", "Platform build (one-off)", F["platform_build"], "S$", note("platform_build"), NUM, True)
    scalar("data_setup", "Data pipelines and cloud set-up (one-off)", F["data_setup"], "S$", note("data_setup"))
    scalar("security_review", "Security and PDPA review (one-off)", F["security_review"], "S$", note("security_review"))
    scalar("training_per_outlet", "Training per outlet (one-off)", F["training_per_outlet"], "S$ per outlet", note("training_per_outlet"))
    scalar("run_team_year", "Run team", F["run_team_year"], "S$ per year", note("run_team_year"))
    scalar("cloud_year", "Cloud, licences and AI services", F["cloud_year"], "S$ per year", note("cloud_year"))
    scalar("maintenance_share", "Software maintenance", F["maintenance_share"], "of platform build a year", note("maintenance_share"), PCT)
    r += 1
    # yearly profile
    A.cell(r, 1, "Yearly profile").font = f(bold=True, color="FFFFFF")
    for c in range(1, 9):
        A.cell(r, c).fill = HEAD
    for i, y in enumerate(yrs):
        A.cell(r, 2 + i, str(y)).font = f(bold=True, color="FFFFFF"); A.cell(r, 2 + i).alignment = Alignment(horizontal="right")
    A.cell(r, 2 + n, "Rationale").font = f(bold=True, color="FFFFFF")
    yr_row_A = r; r += 1
    for key, label, fmt in [("p1_coverage", "Proposal 1: outlets covered", PCT), ("p2_coverage", "Proposal 2: outlets covered", PCT),
                            ("effectiveness", "Model effectiveness (learning curve)", PCT), ("new_stores", "New outlets opened (site model)", NUM),
                            ("team_months", "Project team months", NUM), ("build_share", "One-off build spent", PCT),
                            ("training_outlets", "Outlets trained", NUM), ("run_share", "Run team, share of a year", PCT),
                            ("cloud_share", "Cloud, share of a year", PCT), ("chagee_time_year", "Chagee managers' time (S$)", SGD)]:
        A.cell(r, 1, label).font = f()
        for i, v in enumerate(F[key]):
            c = A.cell(r, 2 + i, v); c.font = f(color=BLUE); c.number_format = fmt
        A.cell(r, 2 + n, note(key)).font = f(color="5C5C5C"); A.merge_cells(start_row=r, start_column=2 + n, end_row=r, end_column=5 + n)
        ref[key] = r; r += 1
    A.column_dimensions["A"].width = 40
    for c in "BCDEF":
        A.column_dimensions[c].width = 13
    A.column_dimensions["C"].width = 16; A.column_dimensions["G"].width = 40; A.column_dimensions["H"].width = 30
    A.freeze_panes = "B4"
    yA = lambda key, i: f"Assumptions!{L(2 + i)}${ref[key]}"            # a yearly input, year i

    # ---- Model
    M = wb.create_sheet("Model")
    M["A1"] = "Chagee Singapore · AI proposals 1-3 · five-year financial impact (S$)"; M["A1"].font = f(bold=True, size=14)
    M["A2"] = "Incremental to Chagee Singapore without the proposals (organic growth and new outlets included in the baseline). Illustrative assumptions, not Chagee data. Execution starts January 2027; the project team is set up in late 2026."
    M["A2"].font = f(italic=True, color="5C5C5C")
    TOT = L(2 + n); NOTE = L(3 + n)
    hdr = 4
    M.cell(hdr, 1, "S$").font = f(bold=True, color="FFFFFF")
    for i, y in enumerate(yrs):
        M.cell(hdr, 2 + i, str(y))
    M.cell(hdr, 2 + n, "5-year total"); M.cell(hdr, 3 + n, "Assumptions")
    for c in range(1, 4 + n):
        M.cell(hdr, c).fill = HEAD; M.cell(hdr, c).font = f(bold=True, color="FFFFFF")
        M.cell(hdr, c).alignment = Alignment(horizontal="left" if c in (1, 3 + n) else "right")
    row = {}; r = hdr + 1

    def head(title):
        nonlocal r
        r += 1
        M.cell(r, 1, title).font = f(bold=True)
        for c in range(1, 4 + n):
            M.cell(r, c).fill = SEC
        r += 1

    def line(key, label, formula, note="", fmt=SGD, total=True, bold=False, border=None, indent=1, color=BLACK):
        """formula(i) -> Excel formula for year column i (0-based)."""
        nonlocal r
        c = M.cell(r, 1, label); c.font = f(bold=bold); c.alignment = Alignment(indent=indent)
        for i in range(n):
            x = M.cell(r, 2 + i, "=" + formula(i)); x.font = f(bold=bold, color=color); x.number_format = fmt
            if border:
                x.border = border
        if total:
            x = M.cell(r, 2 + n, f"=SUM(B{r}:{L(1 + n)}{r})"); x.font = f(bold=True); x.number_format = fmt
            if border:
                x.border = border
        M.cell(r, 3 + n, note).font = f(italic=True, color="5C5C5C", size=9)
        row[key] = r; r += 1

    col = lambda i: L(2 + i)
    at = lambda key, i: f"{col(i)}{row[key]}"
    a = lambda key: ref[key]

    head("Baseline · Chagee Singapore without the proposals")
    line("growth", "Organic growth index", lambda i: f"(1+{a('organic_growth')})^(VALUE({col(i)}${hdr})-VALUE($B${hdr}))", "cups per outlet grow each year", '0.000', total=False)
    line("wage_idx", "Wage index", lambda i: f"(1+{a('wage_inflation')})^(VALUE({col(i)}${hdr})-VALUE($B${hdr}))", "wage inflation; rent held flat", '0.000', total=False)
    line("store_years", "New outlets open (outlet-years)", lambda i: f"SUM(Assumptions!$B${ref['new_stores']}:{L(2 + i)}${ref['new_stores']})-{yA('new_stores', i)}/2",
         "outlets opened mid-year count as half a year", '#,##0.0;(#,##0.0);"-"', total=False)
    line("outlets_avg", "Outlets open (average in the year)", lambda i: f"{a('outlets')}+{at('store_years', i)}", "today's outlets + new outlets", '#,##0.0', total=False)
    line("base_cups", "Cups sold, all outlets", lambda i: f"{at('outlets_avg', i)}*{a('cups')}*{at('growth', i)}", "outlets × cups per outlet × growth", NUM)
    line("base_rev", "Revenue, all outlets", lambda i: f"{at('base_cups', i)}*{a('price')}", "cups × price per cup")

    head("Revenue impact · more cups captured and served")
    line("p1_cups", "Proposal 1 · cups recovered from walk-aways", lambda i: f"{at('base_cups', i)}*{a('p1_lost_share')}*{a('p1_recovery_share')}*{yA('p1_coverage', i)}*{yA('effectiveness', i)}",
         "cups lost at peak × share recovered × coverage × effectiveness", NUM)
    line("p2_cups", "Proposal 2 · extra cups served at the peak", lambda i: f"{at('base_cups', i)}*{a('p2_extra_cup_share')}*{yA('p2_coverage', i)}*{yA('effectiveness', i)}",
         "staff and machines ready for busy days", NUM)
    line("p3_cups", "Proposal 3 · extra cups at better-chosen sites", lambda i: f"{a('cups')}*{a('p3_uplift')}*{at('store_years', i)}*{at('growth', i)}",
         "new outlets × uplift vs a conventional site pick", NUM)
    line("cups", "Additional cups sold", lambda i: f"SUM({col(i)}{row['p1_cups']}:{col(i)}{row['p2_cups']})+{at('p3_cups', i)}", "", NUM, bold=True, border=TOP, indent=0)
    r += 1
    line("p1_rev", "Proposal 1 revenue", lambda i: f"{at('p1_cups', i)}*{a('price')}", "cups × price per cup")
    line("p2_rev", "Proposal 2 revenue", lambda i: f"{at('p2_cups', i)}*{a('price')}", "cups × price per cup")
    line("p3_rev", "Proposal 3 revenue", lambda i: f"{at('p3_cups', i)}*{a('price')}", "cups × price per cup")
    line("rev", "Additional revenue", lambda i: f"SUM({col(i)}{row['p1_rev']}:{col(i)}{row['p3_rev']})", "", bold=True, border=TOP, indent=0)
    line("cogs", "Less ingredients and packaging", lambda i: f"-{at('rev', i)}*(1-{a('margin')})", "revenue × (1 − gross margin)")
    line("gp", "Gross profit from additional sales", lambda i: f"{at('rev', i)}+{at('cogs', i)}", "", bold=True, border=TOP, indent=0)

    head("Cost impact · leaner operations (Proposal 2)")
    line("hours", "Idle staff hours avoided", lambda i: f"{a('p2_idle_hours_week')}*52*{at('outlets_avg', i)}*{yA('p2_coverage', i)}*{yA('effectiveness', i)}",
         "hours a week × 52 × outlets × coverage × effectiveness", NUM)
    line("labour", "Staff overhead saved", lambda i: f"{at('hours', i)}*{a('wage')}*{at('wage_idx', i)}", "hours × labour cost per hour × wage index")
    line("waste", "Fresh stock waste avoided", lambda i: f"{at('base_rev', i)}*{a('p2_waste_share')}*{a('p2_waste_cut')}*{yA('p2_coverage', i)}*{yA('effectiveness', i)}",
         "revenue × waste share × waste avoided")
    line("machines", "Tea machines avoided", lambda i: f"{at('outlets_avg', i)}*{a('p2_machines_saved')}*{a('machine')}*12*{yA('p2_coverage', i)}*{yA('effectiveness', i)}",
         "net machines avoided × monthly cost × 12")
    line("savings", "Operating cost savings", lambda i: f"SUM({col(i)}{row['labour']}:{col(i)}{row['machines']})", "", bold=True, border=TOP, indent=0)
    r += 1
    line("benefit", "TOTAL BENEFIT", lambda i: f"{at('gp', i)}+{at('savings', i)}", "gross profit + cost savings", bold=True, border=TOP, indent=0)

    head("Additional cost from initiatives")
    M.cell(r, 1, "Fixed (one-time)").font = f(italic=True); r += 1
    line("team", "Project team (manager and AI engineer)", lambda i: f"{yA('team_months', i)}*({a('pm_cost_month')}+{a('engineer_cost_month')})", "months × monthly cost; design in 2026, build and pilot in 2027")
    line("platform", "Platform build", lambda i: f"{yA('build_share', i)}*{a('platform_build')}", "external developers: app pop-up and delivery-app link; dashboard on off-the-shelf tools")
    line("data", "Data pipelines and cloud set-up", lambda i: f"{yA('build_share', i)}*{a('data_setup')}", "")
    line("security", "Security and PDPA review", lambda i: f"{yA('build_share', i)}*{a('security_review')}", "before customer data drives app pop-ups")
    line("training", "Training and cross-training", lambda i: f"{yA('training_outlets', i)}*{a('training_per_outlet')}", "by existing zone and store managers: no incremental cost")
    line("fixed", "Total fixed cost", lambda i: f"SUM({col(i)}{row['team']}:{col(i)}{row['training']})", "", bold=True, border=TOP, indent=0)
    M.cell(r, 1, "Variable (recurring)").font = f(italic=True); r += 1
    line("run", "Run team (run and improve the platform)", lambda i: f"{yA('run_share', i)}*{a('run_team_year')}", "from 2028")
    line("maint", "Software maintenance", lambda i: f"{yA('run_share', i)}*{a('maintenance_share')}*{a('platform_build')}", "share of the platform build a year, from 2028")
    line("cloud", "Cloud, licences and AI services", lambda i: f"{yA('cloud_share', i)}*{a('cloud_year')}", "half a year in 2027")
    line("chagee", "Chagee managers' time", lambda i: f"{yA('chagee_time_year', i)}", "existing staff: no incremental cost", color=GREEN)
    line("perks", "Switch perks (Proposal 1)", lambda i: f"{at('base_cups', i)}*{a('p1_redirect_share')}*{yA('p1_coverage', i)}*{yA('effectiveness', i)}*{a('p1_voucher_share')}*{a('p1_voucher_cost')}", "redirected orders × perk share × perk cost")
    line("variable", "Total variable cost", lambda i: f"SUM({col(i)}{row['run']}:{col(i)}{row['perks']})", "", bold=True, border=TOP, indent=0)
    r += 1
    line("cost", "TOTAL COST", lambda i: f"{at('fixed', i)}+{at('variable', i)}", "", bold=True, border=TOP, indent=0)
    r += 1
    line("net", "NET FINANCIAL IMPACT (profit)", lambda i: f"{at('benefit', i)}-{at('cost', i)}", "total benefit − total cost", bold=True, border=DBL, indent=0)
    line("cum", "Cumulative net impact", lambda i: f"SUM($B${row['net']}:{col(i)}{row['net']})", "", bold=True, total=False, indent=0)
    line("paid", "Paid back by this year?", lambda i: f'IF({at("cum", i)}>0,"Yes","No")', "Yes once the cumulative line is above zero", "@", total=False)
    for i in range(n + 1):
        M.cell(row["net"], 2 + i).fill = KEY

    head("Results")
    def kpi(label, formula, fmt, note=""):
        nonlocal r
        M.cell(r, 1, label).font = f(bold=True); M.cell(r, 1).alignment = Alignment(indent=1)
        x = M.cell(r, 2, "=" + formula); x.font = f(bold=True); x.number_format = fmt
        M.cell(r, 3 + n, note).font = f(italic=True, color="5C5C5C", size=9); r += 1
    kpi("Five-year total benefit", f"{TOT}{row['benefit']}", SGD)
    kpi("Five-year total cost", f"{TOT}{row['cost']}", SGD)
    kpi("Five-year net impact", f"{TOT}{row['net']}", SGD)
    kpi("Return on cost (net ÷ cost)", f"IF({TOT}{row['cost']}=0,0,{TOT}{row['net']}/{TOT}{row['cost']})", "0%")
    kpi("Breakeven year (cumulative turns positive)", f"IFERROR(INDEX($B${hdr}:{L(1 + n)}${hdr},MATCH(\"Yes\",$B${row['paid']}:{L(1 + n)}${row['paid']},0)),\"beyond {yrs[-1]}\")", "@",
        "first year the cumulative line is above zero")
    kpi("Full-scale yearly net impact", f"{L(1 + n)}{row['net']}", SGD, f"{yrs[-1]}, all outlets, model fully learned")
    M.column_dimensions["A"].width = 46
    for i in range(n + 1):
        M.column_dimensions[L(2 + i)].width = 13
    M.column_dimensions[NOTE].width = 58
    M.freeze_panes = M.cell(hdr + 1, 2)
    M.cell(row["benefit"], 1).comment = Comment("Gross profit on extra cups plus Proposal 2's operating savings.", "model")

    # ---- SG forecast: the whole of Chagee Singapore, with and without the proposals
    S = wb.create_sheet("SG forecast")
    S["A1"] = "Chagee Singapore · five-year forecast with the three proposals (S$)"; S["A1"].font = f(bold=True, size=14)
    S["A2"] = "Outlet level, before head-office costs and new-store set-up (capex). Illustrative assumptions, not Chagee data."; S["A2"].font = f(italic=True, color="5C5C5C")
    for i, h in enumerate(["S$"] + [str(y) for y in yrs] + ["5-year total", "Notes"], 1):
        x = S.cell(hdr, i, h); x.fill = HEAD; x.font = f(bold=True, color="FFFFFF"); x.alignment = Alignment(horizontal="left" if i in (1, 3 + n) else "right")
    srow = {}; sr = hdr + 1
    def sh(title):
        nonlocal sr
        sr += 1; S.cell(sr, 1, title).font = f(bold=True)
        for c in range(1, 4 + n):
            S.cell(sr, c).fill = SEC
        sr += 1
    def sl(key, label, formula, note="", fmt=SGD, total=True, bold=False, border=None, indent=1, color=BLACK):
        nonlocal sr
        c = S.cell(sr, 1, label); c.font = f(bold=bold); c.alignment = Alignment(indent=indent)
        for i in range(n):
            x = S.cell(sr, 2 + i, "=" + formula(i)); x.font = f(bold=bold, color=color); x.number_format = fmt
            if border:
                x.border = border
        if total:
            x = S.cell(sr, 2 + n, f"=SUM(B{sr}:{L(1 + n)}{sr})"); x.font = f(bold=True); x.number_format = fmt
            if border:
                x.border = border
        S.cell(sr, 3 + n, note).font = f(italic=True, color="5C5C5C", size=9)
        srow[key] = sr; sr += 1
    sat = lambda key, i: f"{col(i)}{srow[key]}"
    mod = lambda key, i: f"Model!{col(i)}{row[key]}"
    sh("Scale")
    sl("outlets_end", "Outlets at year end", lambda i: f"{a('outlets')}+SUM(Assumptions!$B${ref['new_stores']}:{L(2 + i)}${ref['new_stores']})", "today's outlets + new outlets", NUM, total=False)
    sl("cups", "Cups sold", lambda i: f"{mod('base_cups', i)}+{mod('cups', i)}", "baseline + extra cups from the proposals", NUM, color=GREEN)
    sh("Revenue")
    sl("base_rev", "Baseline sales (organic growth and new outlets)", lambda i: mod("base_rev", i), "Model sheet", color=GREEN)
    sl("ai_rev", "Additional sales from the proposals", lambda i: mod("rev", i), "Model sheet", color=GREEN)
    sl("rev", "Total revenue", lambda i: f"{sat('base_rev', i)}+{sat('ai_rev', i)}", "", bold=True, border=TOP, indent=0)
    sl("cogs", "Cost of goods sold", lambda i: f"-{sat('rev', i)}*(1-{a('margin')})+{mod('waste', i)}", "revenue × (1 − gross margin), less fresh-stock waste avoided (Proposal 2)")
    sl("gp", "Gross profit", lambda i: f"{sat('rev', i)}+{sat('cogs', i)}", "", bold=True, border=TOP, indent=0)
    sl("gm", "Gross margin", lambda i: f"IF({sat('rev', i)}=0,0,{sat('gp', i)}/{sat('rev', i)})", "", PCT, total=False)
    sh("Operating expenses")
    sl("labour", "Staff", lambda i: f"-{mod('outlets_avg', i)}*{a('crew_on_shift')}*{a('open_hours')}*{a('open_days')}*{a('wage')}*{mod('wage_idx', i)}+{mod('labour', i)}",
       "outlets × staff on shift × hours × days × wage × wage index, less idle hours saved (Proposal 2)")
    sl("rent", "Rent and utilities", lambda i: f"-{mod('outlets_avg', i)}*({a('rent')}+{a('utilities')})*12", "outlets × monthly cost × 12")
    sl("delivery", "Delivery-app commissions", lambda i: f"-{sat('base_rev', i)}*{a('delivery_share')}*{a('delivery_commission')}", "baseline revenue × delivery share × commission")
    sl("equip", "Equipment: tea machines avoided", lambda i: mod("machines", i), "Proposal 2, Model sheet", color=GREEN)
    sl("ai_cost", "AI platform and team", lambda i: f"-{mod('cost', i)}", "cost of the initiatives, Model sheet", color=GREEN)
    sl("opex", "Total operating expenses", lambda i: f"SUM({col(i)}{srow['labour']}:{col(i)}{srow['ai_cost']})", "", bold=True, border=TOP, indent=0)
    r_ = sr; sr += 1
    sl("op", "OUTLET OPERATING PROFIT (before head-office costs)", lambda i: f"{sat('gp', i)}+{sat('opex', i)}", "", bold=True, border=DBL, indent=0)
    sl("margin", "Operating margin", lambda i: f"IF({sat('rev', i)}=0,0,{sat('op', i)}/{sat('rev', i)})", "", PCT, total=False)
    for i in range(n + 1):
        S.cell(srow["op"], 2 + i).fill = KEY
    sh("Memo: without the proposals")
    sl("op_base", "Outlet operating profit without the proposals", lambda i: f"{sat('base_rev', i)}*{a('margin')}-{mod('outlets_avg', i)}*{a('crew_on_shift')}*{a('open_hours')}*{a('open_days')}*{a('wage')}*{mod('wage_idx', i)}+{sat('rent', i)}+{sat('delivery', i)}",
       "baseline gross profit − staff − rent and utilities − delivery commissions")
    sl("margin_base", "Operating margin without the proposals", lambda i: f"IF({sat('base_rev', i)}=0,0,{sat('op_base', i)}/{sat('base_rev', i)})", "", PCT, total=False)
    sl("uplift", "Uplift from the proposals", lambda i: f"{sat('op', i)}-{sat('op_base', i)}", "equals the Model sheet's net impact", bold=True, border=TOP, indent=0)
    S.column_dimensions["A"].width = 48
    for i in range(n + 1):
        S.column_dimensions[L(2 + i)].width = 13
    S.column_dimensions[NOTE].width = 46
    S.freeze_panes = S.cell(hdr + 1, 2)

    # ---- Cost plan (by stage, like a project budget)
    C = wb.create_sheet("Cost plan")
    C["A1"] = "Cost of the initiatives by stage (S$)"; C["A1"].font = f(bold=True, size=14)
    C["A2"] = "Linked to the Assumptions sheet. Fixed costs are one-time; recurring costs are per year at full scale."; C["A2"].font = f(italic=True, color="5C5C5C")
    cols = ["Stage", "Timing", "Task", "Resources", "Fixed, one-time", "Recurring, per year"]
    for i, h in enumerate(cols, 1):
        x = C.cell(4, i, h); x.font = f(bold=True, color="FFFFFF"); x.fill = HEAD
    pm, en = a("pm_cost_month"), a("engineer_cost_month")
    plan = [
        ("Design and plan", "Oct–Dec 2026", "Design the three proposals, audit data, plan the pilots", "Project manager, AI engineer", f"={yA('team_months', 0)}*({pm}+{en})", None),
        ("Build and pilot", "2027", "Build the platform: app pop-up and delivery-app link; configure the dashboard and assistant on off-the-shelf tools", "External developers", f"={a('platform_build')}", None),
        ("", "2027", "Data pipelines and cloud set-up", "AI engineer", f"={a('data_setup')}", None),
        ("", "2027", "Security testing and PDPA review", "External specialist", f"={a('security_review')}", None),
        ("", "2027", "Run the project: pilots, go/no-go, roll-out", "Project manager, AI engineer", f"={yA('team_months', 1)}*({pm}+{en})", None),
        ("", "2027", "Train and cross-train outlet teams", "Zone and store managers (existing staff)", f"={yA('training_outlets', 1)}*{a('training_per_outlet')}", None),
        ("Run and scale", "2028 onwards", "Run and improve the platform", "Run team", None, f"={a('run_team_year')}"),
        ("", "2028 onwards", "Maintain the externally built software", "External developers", None, f"={a('maintenance_share')}*{a('platform_build')}"),
        ("", "2028 onwards", "Cloud, licences and AI services", "Cloud provider", None, f"={a('cloud_year')}"),
        ("", "2028 onwards", "Chagee managers' time", "Supply chain, zone and store managers (existing staff)", None, f"={yA('chagee_time_year', 2)}"),
        ("", "2028 onwards", "Switch perks for redirected customers", "Proposal 1", None, f"=Model!{L(1 + n)}{row['perks']}"),
        ("", "2028 onwards", "Train each new outlet", "Store manager (existing staff)", None, f"={yA('training_outlets', 2)}*{a('training_per_outlet')}"),
    ]
    rr = 5
    for st, tm, task, res, fx, rc in plan:
        for i, v in enumerate((st, tm, task, res), 1):
            C.cell(rr, i, v).font = f(bold=bool(st) and i == 1)
            C.cell(rr, i).alignment = Alignment(wrap_text=True, vertical="top")
        for i, v in ((5, fx), (6, rc)):
            if v:
                x = C.cell(rr, i, v); x.font = f(); x.number_format = SGD; x.alignment = Alignment(vertical="top")
        rr += 1
    C.cell(rr, 1, "Total").font = f(bold=True)
    for i in (5, 6):
        x = C.cell(rr, i, f"=SUM({L(i)}5:{L(i)}{rr - 1})"); x.font = f(bold=True); x.number_format = SGD; x.border = DBL
    for c, w in zip("ABCDEF", (18, 14, 56, 30, 16, 18)):
        C.column_dimensions[c].width = w

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb.active = 1
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)      # openpyxl stores no results: Excel computes them on open
    wb.save(path)
    return path


if __name__ == "__main__":
    out = export()
    res = compute(); R, T = res["rows"], res["total"]
    for k in ("rev", "gp", "savings", "benefit", "fixed", "variable", "cost", "net", "cum"):
        print(f"{k:9}", " ".join(f"{x:>12,.0f}" for x in R[k]), f"| {T.get(k, 0):>12,.0f}")
    print("roi", f"{T['roi']:.0%}", "payback", T["payback"], "->", out)
